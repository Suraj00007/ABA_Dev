"""Build an Excel workbook documenting every filter, grouping, column and setting of the ABA
reports, dashboards and Home page, straight from the deployed metadata, plus live opportunity data
with per-report membership flags.

Usage (from the project root):
    python scripts/aba/build_filter_workbook.py
"""
import datetime as dt
import json
import os
import subprocess
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SRC = os.path.join(ROOT, "force-app", "main", "default")
OUT = os.path.join(ROOT, "ABA_Reports_Dashboards_Filter_Criteria.xlsx")
NS = {"m": "http://soap.sforce.com/2006/04/metadata"}
ORG = "ABA_Dev"
AS_OF = dt.date(2026, 10, 8)

# --------------------------------------------------------------------------- labels
FIELD_LABELS = {
    "OPPORTUNITY_NAME": "Opportunity Name", "ACCOUNT_NAME": "Account Name", "FULL_NAME": "Opportunity Owner",
    "AMOUNT": "Amount", "CLOSE_DATE": "Close Date", "CLOSE_MONTH": "Close Month", "STAGE_NAME": "Stage",
    "TYPE": "Type", "PROBABILITY": "Probability (%)", "EXP_AMOUNT": "Expected Revenue", "NEXT_STEP": "Next Step",
    "FORECAST_CATEGORY": "Forecast Category", "WON": "Won", "CLOSED": "Closed", "LAST_ACTIVITY": "Last Activity",
    "LAST_UPDATE": "Last Modified Date", "AGE": "Age", "CREATED_DATE": "Created Date", "LEAD_SOURCE": "Lead Source",
    "CAMPAIGN_SOURCE": "Primary Campaign Source", "INDUSTRY": "Account: Industry",
    "ADDRESS1_COUNTRY_CODE": "Account: Billing Country", "Account.Segment__c": "Account: Segment",
    "Opportunity.Win_Loss_Reason__c": "Win/Loss Reason", "Opportunity.Win_Loss_Details__c": "Win/Loss Details",
    "Opportunity.Proposal_Submission_Date__c": "Proposal Submission Date",
    "Opportunity.Deferred_Date__c": "Deferred Date", "Opportunity.Close_Quarter__c": "Close Quarter",
    "FORMULA1": "Win Rate % (custom summary formula)", "RowCount": "Record Count", "s!AMOUNT": "Sum of Amount",
}
INTERVALS = {
    "INTERVAL_CURFY": ("Current FY", "1 Jan 2026 – 31 Dec 2026"),
    "INTERVAL_THISMONTH": ("Current Calendar Month", "1 Oct 2026 – 31 Oct 2026"),
    "INTERVAL_CURRENT": ("Current FQ", "1 Oct 2026 – 31 Dec 2026"),
    "INTERVAL_PREVCURFY": ("Previous FY + Current FY", "1 Jan 2025 – 31 Dec 2026"),
    "INTERVAL_CURNEXT3": ("Current FQ + Next 3 FQ", "1 Oct 2026 – 30 Sep 2027"),
    "INTERVAL_NEXT30": ("Next 30 Days", "Rolling: the next 30 days from today"),
    "INTERVAL_CUSTOM": ("All Time", "No date limit (custom range left empty)"),
}
OPEN_PARAM = {"all": "All Opportunities", "open": "Open Opportunities", "closed": "Closed Opportunities",
              "closedwon": "Closed Won Opportunities"}
OPERATORS = {"equals": "equals", "notEqual": "not equal to", "lessThan": "less than", "greaterThan": "greater than",
             "lessOrEqual": "less or equal", "greaterOrEqual": "greater or equal"}
GRANULARITY = {"None": "", "FiscalQuarter": "Fiscal Quarter", "FiscalYear": "Fiscal Year", "Month": "Calendar Month"}
FOLDER_SHARING = {"ABA_Management_Reports": "View: role CEO and subordinates (internal)",
                  "ABA_Sales_Reports": "View: all internal users",
                  "ABA_Management_Dashboards": "View: role CEO and subordinates (internal)"}

# Purpose / management question per report (from the build guide)
PURPOSE = {
    "R01": ("How much revenue have we won this year?", "Q1"),
    "R02": ("How much revenue have we won this month?", "Q1"),
    "R03": ("How much revenue have we won this quarter, against target?", "Q1"),
    "R04": ("How does revenue compare quarter by quarter, this year vs last year?", "Q2"),
    "R05": ("What does the active pipeline look like by stage?", "Q3"),
    "R06": ("Which quarter is the pipeline expected to close in, by forecast category?", "Q3"),
    "R07": ("How much pipeline does each salesperson carry, by stage?", "Q4"),
    "R08": ("Who has won the most revenue this year (leaderboard)?", "Q4"),
    "R09": ("What is each salesperson's win rate this year?", "Q4"),
    "R10": ("Why do we win and why do we lose?", "Q5"),
    "R10a": ("Why do we win? (dashboard source, won deals only)", "Q5"),
    "R10b": ("Why do we lose? (dashboard source, lost deals only)", "Q5"),
    "R11": ("Which deals at Proposal Submitted or later close in the next 30 days?", "Q6"),
    "R12": ("Which open deals are overdue or have had no activity for 30 days?", "Q7"),
    "R13": ("Which deals are deferred and what value could come back?", "Q8"),
    "R14": ("How much comes from new logos vs existing customers?", "Q9"),
    "R15": ("Which lead sources and campaigns generate revenue?", "Q10"),
    "R16": ("Which industries, segments and countries bring revenue?", "Q11"),
    "R17": ("Which proposals are due in the next 30 days?", "Q12"),
    "R18": ("What are our average deal size and sales cycle?", "Q13"),
}

# SOQL equivalent of each report's filters - used to flag which opportunities are in which report.
ACTIVE = "IsClosed = false AND StageName NOT IN ('Deferred','Old / Archived')"
SOQL_WHERE = {
    "R01": "IsWon = true AND StageName = 'Closed Won' AND CloseDate = THIS_FISCAL_YEAR",
    "R02": "StageName = 'Closed Won' AND CloseDate = THIS_MONTH",
    "R03": "StageName = 'Closed Won' AND CloseDate = THIS_FISCAL_QUARTER",
    "R04": "StageName = 'Closed Won' AND (CloseDate = THIS_FISCAL_YEAR OR CloseDate = LAST_FISCAL_YEAR)",
    "R05": ACTIVE,
    "R06": ACTIVE + " AND (CloseDate = THIS_FISCAL_QUARTER OR CloseDate = NEXT_N_FISCAL_QUARTERS:3)",
    "R07": ACTIVE,
    "R08": "StageName = 'Closed Won' AND CloseDate = THIS_FISCAL_YEAR",
    "R09": "IsClosed = true AND CloseDate = THIS_FISCAL_YEAR",
    "R10": "IsClosed = true AND CloseDate = THIS_FISCAL_YEAR",
    "R10a": "StageName = 'Closed Won' AND CloseDate = THIS_FISCAL_YEAR",
    "R10b": "StageName = 'Closed Lost' AND CloseDate = THIS_FISCAL_YEAR",
    "R11": "IsClosed = false AND StageName IN ('Proposal Submitted','Negotiation','Verbal Award') AND CloseDate = NEXT_N_DAYS:30",
    "R12": ACTIVE + " AND (CloseDate < TODAY OR LastActivityDate < LAST_N_DAYS:30)",
    "R13": "StageName = 'Deferred'",
    "R14": "CloseDate = THIS_FISCAL_YEAR AND StageName NOT IN ('Closed Lost','Deferred','Old / Archived')",
    "R15": "StageName = 'Closed Won' AND CloseDate = THIS_FISCAL_YEAR",
    "R16": "StageName = 'Closed Won' AND CloseDate = THIS_FISCAL_YEAR",
    "R17": "IsClosed = false AND Proposal_Submission_Date__c = NEXT_N_DAYS:30",
    "R18": "StageName = 'Closed Won' AND CloseDate = LAST_N_DAYS:365",
}


def label(api):
    return FIELD_LABELS.get(api, api)


# --------------------------------------------------------------------------- Salesforce access
def sf(*args):
    out = subprocess.run(["sf.cmd", *args, "-o", ORG], capture_output=True, text=True, encoding="utf-8").stdout
    return out


def soql(q):
    data = json.loads(sf("data", "query", "-q", q, "--json"))
    if data.get("status") != 0:
        raise RuntimeError(data.get("message"))
    return data["result"]["records"]


# --------------------------------------------------------------------------- metadata parsing
def txt(el, path):
    found = el.find(path, NS)
    return found.text if found is not None else None


def parse_report(path, folder):
    root = ET.parse(path).getroot()
    name = txt(root, "m:name")
    rid = name.split(" ")[0]
    r = {"id": rid, "name": name[len(rid) + 1:], "api": os.path.basename(path).replace(".report-meta.xml", ""),
         "folder": folder, "format": txt(root, "m:format"), "type": txt(root, "m:reportType"),
         "scope": txt(root, "m:scope"), "description": txt(root, "m:description"),
         "row_limit": txt(root, "m:rowLimit"), "sort_col": txt(root, "m:sortColumn"), "sort_order": txt(root, "m:sortOrder"),
         "show_details": txt(root, "m:showDetails"), "grand_total": txt(root, "m:showGrandTotal"),
         "subtotals": txt(root, "m:showSubTotals")}
    r["params"] = {txt(p, "m:name"): txt(p, "m:value") for p in root.findall("m:params", NS)}
    tf = root.find("m:timeFrameFilter", NS)
    r["date_col"], r["interval"] = txt(tf, "m:dateColumn"), txt(tf, "m:interval")
    flt = root.find("m:filter", NS)
    r["logic"] = txt(flt, "m:booleanFilter") if flt is not None else None
    r["criteria"] = [(txt(c, "m:column"), txt(c, "m:operator"), txt(c, "m:value"))
                     for c in (flt.findall("m:criteriaItems", NS) if flt is not None else [])]
    r["columns"] = [(txt(c, "m:field"), ", ".join(a.text for a in c.findall("m:aggregateTypes", NS)))
                    for c in root.findall("m:columns", NS)]

    def groups(tagname):
        return [{"field": txt(g, "m:field"), "gran": txt(g, "m:dateGranularity"), "order": txt(g, "m:sortOrder"),
                 "sort_by": txt(g, "m:sortByName"), "sort_type": txt(g, "m:sortType")}
                for g in root.findall(tagname, NS)]
    r["down"], r["across"] = groups("m:groupingsDown"), groups("m:groupingsAcross")
    agg = root.find("m:aggregates", NS)
    r["formula"] = None if agg is None else {
        "label": txt(agg, "m:masterLabel"), "api": txt(agg, "m:developerName"), "calc": txt(agg, "m:calculatedFormula"),
        "type": txt(agg, "m:datatype"), "scale": txt(agg, "m:scale")}
    ch = root.find("m:chart", NS)
    r["chart"] = None if ch is None else {
        "type": txt(ch, "m:chartType"), "group": txt(ch, "m:groupingColumn"), "secondary": txt(ch, "m:secondaryGroupingColumn"),
        "summary": f"{txt(ch, 'm:chartSummaries/m:aggregate') or ''} of {label(txt(ch, 'm:chartSummaries/m:column'))}".strip(),
        "legend": txt(ch, "m:legendPosition"), "values": txt(ch, "m:showValues"), "pct": txt(ch, "m:showPercentage"),
        "total": txt(ch, "m:showTotal"), "size": txt(ch, "m:size")}
    return r


def parse_dashboard(path):
    root = ET.parse(path).getroot()
    d = {"api": os.path.basename(path).replace(".dashboard-meta.xml", ""), "title": txt(root, "m:title"),
         "description": txt(root, "m:description"), "type": txt(root, "m:dashboardType"),
         "running_user": txt(root, "m:runningUser"), "grid_cols": txt(root, "m:dashboardGridLayout/m:numberOfColumns"),
         "row_height": txt(root, "m:dashboardGridLayout/m:rowHeight"), "palette": txt(root, "m:dashboardColorPalette"),
         "theme": txt(root, "m:dashboardChartTheme")}
    d["filters"] = [(txt(f, "m:name"), [txt(o, "m:values") for o in f.findall("m:dashboardFilterOptions", NS)],
                     [txt(o, "m:operator") for o in f.findall("m:dashboardFilterOptions", NS)])
                    for f in root.findall("m:dashboardFilters", NS)]
    comps = []
    for g in root.findall("m:dashboardGridLayout/m:dashboardGridComponents", NS):
        c = g.find("m:dashboardComponent", NS)
        flex = c.find("m:flexComponentProperties", NS)
        comps.append({
            "row": int(txt(g, "m:rowIndex")), "col": int(txt(g, "m:columnIndex")),
            "rows": int(txt(g, "m:rowSpan")), "cols": int(txt(g, "m:colSpan")),
            "header": txt(c, "m:header"), "type": txt(c, "m:componentType"), "report": txt(c, "m:report"),
            "summary": (txt(c, "m:chartSummary/m:aggregate"), txt(c, "m:chartSummary/m:column")),
            "groupings": [e.text for e in c.findall("m:groupingColumn", NS)],
            "sort_by": txt(c, "m:sortBy"), "legend": txt(c, "m:legendPosition"),
            "show_values": txt(c, "m:showValues"), "show_total": txt(c, "m:showTotal"), "show_pct": txt(c, "m:showPercentage"),
            "display_units": txt(c, "m:displayUnits"), "metric_label": txt(c, "m:metricLabel"),
            "gauge": (txt(c, "m:gaugeMin"), txt(c, "m:gaugeMax"), txt(c, "m:indicatorBreakpoint1"), txt(c, "m:indicatorBreakpoint2")),
            "colors": (txt(c, "m:indicatorLowColor"), txt(c, "m:indicatorMiddleColor"), txt(c, "m:indicatorHighColor")),
            "filter_cols": [e.text for e in c.findall("m:dashboardFilterColumns/m:column", NS)],
            "table_cols": [(txt(t, "m:reportColumn"), txt(t, "m:type")) for t in flex.findall("m:flexTableColumn", NS)] if flex is not None else [],
            "table_sort": (txt(flex, "m:flexTableSortInfo/m:sortColumn"), txt(flex, "m:flexTableSortInfo/m:sortOrder")) if flex is not None else None,
            "drill": txt(c, "m:drillEnabled"), "auto": txt(c, "m:autoselectColumnsFromReport"),
        })
    d["components"] = sorted(comps, key=lambda c: (c["row"], c["col"]))
    return d


def parse_flexipage(path):
    root = ET.parse(path).getroot()
    rows = []
    for reg in root.findall("m:flexiPageRegions", NS):
        rname = txt(reg, "m:name")
        for i, inst in enumerate(reg.findall("m:itemInstances/m:componentInstance", NS), 1):
            props = {txt(p, "m:name"): txt(p, "m:value") for p in inst.findall("m:componentInstanceProperties", NS)}
            rows.append((rname, i, txt(inst, "m:componentName"), props))
    return txt(root, "m:masterLabel"), txt(root, "m:template/m:name"), rows


# --------------------------------------------------------------------------- styling helpers
ARIAL = "Arial"
HDR_FILL = PatternFill("solid", start_color="1F4E5A")
SUB_FILL = PatternFill("solid", start_color="DCEBEE")
NOTE_FILL = PatternFill("solid", start_color="FFF4D6")
THIN = Side(style="thin", color="C9D3D8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_sheet(ws, widths, header_row=1, freeze="B2", table_name=None):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=header_row):
        for cell in row:
            if cell.font.name != ARIAL or not cell.font.b:
                cell.font = Font(name=ARIAL, size=10, bold=cell.font.b, color=cell.font.color)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
    for cell in ws[header_row]:
        cell.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
        cell.fill = HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[header_row].height = 30
    if freeze:
        ws.freeze_panes = freeze
    if table_name and ws.max_row > header_row:
        ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
        t = Table(displayName=table_name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(t)


def title_block(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = Font(name=ARIAL, size=14, bold=True, color="1F4E5A")
    ws["A2"] = subtitle
    ws["A2"].font = Font(name=ARIAL, size=10, italic=True, color="566572")


# --------------------------------------------------------------------------- build
def main():
    reports = []
    for folder in ("ABA_Management_Reports", "ABA_Sales_Reports"):
        d = os.path.join(SRC, "reports", folder)
        for f in sorted(os.listdir(d)):
            reports.append(parse_report(os.path.join(d, f), folder))
    order = list(PURPOSE)
    reports.sort(key=lambda r: order.index(r["id"]))
    by_api = {f"{r['folder']}/{r['api']}": r for r in reports}

    dash_dir = os.path.join(SRC, "dashboards", "ABA_Management_Dashboards")
    dashboards = [parse_dashboard(os.path.join(dash_dir, f)) for f in sorted(os.listdir(dash_dir))]
    dashboards.sort(key=lambda d: 0 if "Management" in d["api"] else 1)
    for i, d in enumerate(dashboards, 1):
        d["id"] = f"D{i}"
    home_label, home_template, home_rows = parse_flexipage(os.path.join(SRC, "flexipages", "ABA_Sales_Home_Page.flexipage-meta.xml"))

    # which dashboard components / home regions use each report
    used_in = {r["id"]: [] for r in reports}
    for d in dashboards:
        for c in d["components"]:
            used_in[by_api[c["report"]]["id"]].append(f"{d['id']}: {c['header']}")
    rep_by_dev = {r["api"]: r for r in reports}
    for region, _, comp, props in home_rows:
        if comp == "flexipage:reportChart":
            used_in[rep_by_dev[props["reportName"]]["id"]].append(f"Home page ({region}): report chart")

    # live data
    opps = soql("SELECT Id, Name, Account.Name, Account.Industry, Account.BillingCountry, Account.Segment__c, Owner.Name, "
                "StageName, IsClosed, IsWon, Probability, Amount, CloseDate, Close_Quarter__c, FiscalYear, ForecastCategoryName, "
                "Type, LeadSource, Campaign.Name, Offering__c, Proposal_Submission_Date__c, Win_Loss_Reason__c, Win_Loss_Details__c, "
                "Deferred_Date__c, NextStep, LastActivityDate, CreatedDate FROM Opportunity ORDER BY CloseDate, Name")
    member = {}
    for rid, where in SOQL_WHERE.items():
        member[rid] = {r["Id"] for r in soql(f"SELECT Id FROM Opportunity WHERE {where}")}
    live = {}
    for r in soql("SELECT Id, DeveloperName FROM Report WHERE FolderName IN ('ABA Management Reports','ABA Sales Reports')"):
        res = json.loads(sf("api", "request", "rest", f"/services/data/v67.0/analytics/reports/{r['Id']}?includeDetails=true"))  # row-limited reports misreport counts without details
        fm = res.get("factMap", {})
        key = "T!T" if "T!T" in fm else next((k for k in fm if k.endswith("T!T")), None)
        aggs = fm.get(key, {}).get("aggregates", []) if key else []
        live[rep_by_dev[r["DeveloperName"]]["id"]] = (r["Id"], aggs[-1]["value"] if aggs else None)

    wb = Workbook()

    # ------------------------------------------------------------- 1. Overview
    ws = wb.active
    ws.title = "Overview"
    title_block(ws, "ABA Opportunity Reports & Dashboards – Filter Criteria",
                "Generated from the deployed Salesforce metadata and live data in org ABA_Dev.")
    rows = [
        ("As of date", AS_OF, "Date the live data was read. Relative date filters (this month, current FY, next 30 days) resolve against this date. Change it to recalculate Age on the Opportunity Data sheet."),
        ("Org", "ABA_Dev (00DgK00000bWvvVUAS)", "Developer Edition"),
        ("Fiscal year", "Standard, January – December", "Fiscal quarters equal calendar quarters"),
        ("Currency", "Amounts are SAR values from the sample CSV", "Org currency is USD, so Salesforce shows a $ sign"),
        ("Report type", "Opportunities (standard)", "All reports use Show Me = All opportunities (scope: organization)"),
    ]
    ws.append([])
    ws.append(["Setting", "Value", "Note"])
    for r in rows:
        ws.append(list(r))
    ws["B5"].number_format = "dd mmm yyyy"
    ws["B5"].font = Font(name=ARIAL, size=10, color="0000FF")
    ws["B5"].fill = PatternFill("solid", start_color="FFFF00")
    ws["B5"].comment = Comment("Input: the date the data was extracted from Salesforce (8 Oct 2026). Used by the Age formula.", "ABA build")
    ws.append([])
    ws.append(["Sheet", "What it shows", ""])
    sheets = [
        ("Report Catalog", "One row per report: purpose, folder, format, every filter, grouping, column, summary, chart and sort, and where it is used."),
        ("Report Filters", "One row per filter condition per report, including the standard Opportunity Status and date-range filters and the filter logic."),
        ("Report Fields", "One row per field per report, and its role (filter, date filter, row or column grouping, detail column, summary, chart)."),
        ("Field Usage Matrix", "Every field against every report, showing where each field is used."),
        ("Dashboards", "Dashboard-level settings: folder, running user, grid, and dashboard filters with all their options."),
        ("Dashboard Components", "One row per component: position on the grid, type, source report and that report's filters, measure, groupings, sort, gauge bands, table columns and current value."),
        ("Home Page", "Sales app Home page regions and their components."),
        ("Opportunity Data", "Every opportunity in the org with all reporting fields, plus a 1/0 flag showing whether it is included in each report."),
        ("Report Results", "Record count and amount for each report, calculated with formulas from the flags, alongside the live count from Salesforce."),
        ("Folders & Access", "Report and dashboard folders and who they are shared with."),
    ]
    for s in sheets:
        ws.append([s[0], s[1], ""])
    ws.append([])
    ws.append(["Legend", "Yellow fill and blue text = input you can change. Black text = formula or text taken from Salesforce.", ""])
    style_sheet(ws, [26, 70, 70], header_row=4, freeze=None)
    for cell in ws[11]:
        cell.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
        cell.fill = HDR_FILL
    ws["B5"].font = Font(name=ARIAL, size=10, color="0000FF")
    ws["A1"].font = Font(name=ARIAL, size=14, bold=True, color="1F4E5A")
    ws["A2"].font = Font(name=ARIAL, size=10, italic=True, color="566572")
    for c in ("A1", "A2"):
        ws[c].border = Border()
    as_of_ref = "Overview!$B$5"

    # ------------------------------------------------------------- 2. Report Catalog
    ws = wb.create_sheet("Report Catalog")
    hdr = ["Report ID", "Report name", "API name", "Folder", "Management question", "Guide Q#", "Description",
           "Format", "Report type", "Show Me (scope)", "Opportunity Status", "Date field", "Date range",
           "Date range resolved (as of 8 Oct 2026)", "Field filters", "Filter logic", "Group rows by", "Group columns by",
           "Detail columns", "Summaries", "Custom summary formula", "Row limit", "Sort", "Report chart",
           "Show details / subtotals / grand total", "Used in dashboards & Home page", "Salesforce report ID"]
    ws.append(hdr)
    for r in reports:
        crit = "\n".join(f"{i}. {label(c)} {OPERATORS.get(o, o)} \"{v}\"" for i, (c, o, v) in enumerate(r["criteria"], 1)) or "None"
        down = " → ".join(label(g["field"]) + (f" ({GRANULARITY[g['gran']]})" if GRANULARITY.get(g["gran"]) else "")
                          + (f", sorted by {g['sort_type'].lower()} {label(g['sort_by'])} {g['order'].lower()}" if g["sort_by"] else "")
                          for g in r["down"]) or "–"
        across = " → ".join(label(g["field"]) + (f" ({GRANULARITY[g['gran']]})" if GRANULARITY.get(g["gran"]) else "")
                            for g in r["across"]) or "–"
        detail = ", ".join(label(c) for c, _ in r["columns"])
        summ = ", ".join(f"{a} of {label(c)}" for c, a in r["columns"] if a) + ", Record Count"
        formula = (f"{r['formula']['label']} = {r['formula']['calc']} ({r['formula']['type']}, {r['formula']['scale']} dp)"
                   if r["formula"] else "–")
        sort = f"{label(r['sort_col'])} {r['sort_order']}" if r["sort_col"] else "Default"
        ch = r["chart"]
        chart = (f"{ch['type']} · X: {label(ch['group'])}" + (f" · Series: {label(ch['secondary'])}" if ch["secondary"] else "")
                 + f" · Y: {ch['summary']}" + (f" · Legend {ch['legend']}" if ch["legend"] else "")) if ch else "No chart"
        interval = INTERVALS[r["interval"]]
        ws.append([r["id"], r["name"], r["api"], r["folder"].replace("_", " "), PURPOSE[r["id"]][0], PURPOSE[r["id"]][1],
                   r["description"], r["format"], r["type"], "All opportunities" if r["scope"] == "organization" else r["scope"],
                   OPEN_PARAM.get(r["params"].get("open"), r["params"].get("open")), label(r["date_col"]), interval[0],
                   interval[1], crit, r["logic"] or "AND (all conditions)", down, across, detail, summ, formula,
                   r["row_limit"] or "None", sort, chart,
                   f"{r['show_details']} / {r['subtotals']} / {r['grand_total']}", "\n".join(used_in[r["id"]]) or "Report only",
                   live.get(r["id"], ("", None))[0]])
    style_sheet(ws, [9, 34, 34, 22, 40, 8, 40, 10, 12, 16, 20, 20, 22, 28, 48, 18, 40, 24, 50, 30, 36, 9, 22, 46, 18, 44, 20],
                freeze="C2", table_name="ReportCatalog")

    # ------------------------------------------------------------- 3. Report Filters
    ws = wb.create_sheet("Report Filters")
    ws.append(["Report ID", "Report name", "Filter #", "Filter kind", "Field", "Field API name", "Operator", "Value",
               "Resolved value (as of 8 Oct 2026)", "Filter logic", "Plain-English meaning"])
    for r in reports:
        ws.append([r["id"], r["name"], "S1", "Standard filter: Show Me", "Scope", "scope", "equals", "All opportunities",
                   "All opportunities", "", "Every user's opportunities are included, not just the viewer's."])
        status = r["params"].get("open")
        ws.append([r["id"], r["name"], "S2", "Standard filter: Opportunity Status", "Opportunity Status", "open", "equals",
                   OPEN_PARAM.get(status, status), OPEN_PARAM.get(status, status), "",
                   {"all": "Open and closed opportunities.", "open": "Only opportunities that are not closed.",
                    "closed": "Only Closed Won and Closed Lost opportunities."}.get(status, "")])
        interval = INTERVALS[r["interval"]]
        ws.append([r["id"], r["name"], "S3", "Standard filter: Date range", label(r["date_col"]), r["date_col"], "within",
                   interval[0], interval[1], "", f"{label(r['date_col'])} falls in {interval[0].lower()}."])
        for i, (c, o, v) in enumerate(r["criteria"], 1):
            values = v.split(",")
            meaning = (f"{label(c)} is " + ("not " if o == "notEqual" else "") + (" or ".join(values) if o in ("equals", "notEqual") else f"{OPERATORS[o]} {v}") + ".")
            resolved = {"TODAY": "8 Oct 2026", "LAST_N_DAYS:30": "before 8 Sep 2026", "LAST_N_DAYS:365": "8 Oct 2025 – 8 Oct 2026"}.get(v, v)
            ws.append([r["id"], r["name"], i, "Field filter", label(c), c, OPERATORS.get(o, o), v, resolved,
                       r["logic"] or "AND", meaning])
    style_sheet(ws, [9, 34, 8, 26, 24, 34, 14, 36, 30, 14, 52], freeze="C2", table_name="ReportFilters")

    # ------------------------------------------------------------- 4. Report Fields
    ws = wb.create_sheet("Report Fields")
    ws.append(["Report ID", "Report name", "Role", "Order", "Field", "Field API name", "Aggregate / granularity", "Sort"])
    usage = {}

    def use(field, rid, role):
        usage.setdefault(field, {}).setdefault(rid, [])
        if role not in usage[field][rid]:
            usage[field][rid].append(role)

    for r in reports:
        rid = r["id"]
        for c, o, v in r["criteria"]:
            ws.append([rid, r["name"], "Field filter", "", label(c), c, f"{OPERATORS.get(o, o)} {v}", ""])
            use(c, rid, "F")
        ws.append([rid, r["name"], "Date filter", "", label(r["date_col"]), r["date_col"], INTERVALS[r["interval"]][0], ""])
        use(r["date_col"], rid, "D")
        for i, g in enumerate(r["down"], 1):
            sort = f"{g['sort_type']} {label(g['sort_by'])} {g['order']}" if g["sort_by"] else g["order"]
            ws.append([rid, r["name"], "Group rows by", i, label(g["field"]), g["field"], GRANULARITY.get(g["gran"], g["gran"]), sort])
            use(g["field"], rid, "GR")
        for i, g in enumerate(r["across"], 1):
            ws.append([rid, r["name"], "Group columns by", i, label(g["field"]), g["field"], GRANULARITY.get(g["gran"], g["gran"]), g["order"]])
            use(g["field"], rid, "GC")
        for i, (c, a) in enumerate(r["columns"], 1):
            ws.append([rid, r["name"], "Detail column", i, label(c), c, a or "", ""])
            use(c, rid, "C" + ("Σ" if a else ""))
        if r["formula"]:
            ws.append([rid, r["name"], "Custom summary formula", "", r["formula"]["label"], r["formula"]["api"], r["formula"]["calc"], ""])
            use("FORMULA1", rid, "SF")
        if r["sort_col"]:
            ws.append([rid, r["name"], "Report sort", "", label(r["sort_col"]), r["sort_col"], "", r["sort_order"]])
            use(r["sort_col"], rid, "S")
        if r["chart"]:
            ws.append([rid, r["name"], "Chart X-axis", "", label(r["chart"]["group"]), r["chart"]["group"], r["chart"]["type"], ""])
            if r["chart"]["secondary"]:
                ws.append([rid, r["name"], "Chart series", "", label(r["chart"]["secondary"]), r["chart"]["secondary"], r["chart"]["type"], ""])
    style_sheet(ws, [9, 34, 22, 7, 30, 38, 30, 34], freeze="C2", table_name="ReportFields")

    # ------------------------------------------------------------- 5. Field Usage Matrix
    ws = wb.create_sheet("Field Usage Matrix")
    ids = [r["id"] for r in reports]
    ws.append(["Field", "Field API name"] + ids + ["Reports using it"])
    for field in sorted(usage, key=lambda f: label(f).lower()):
        row = [label(field), field] + [" ".join(usage[field].get(i, [])) for i in ids]
        ws.append(row)
        n = ws.max_row
        ws.cell(n, len(row) + 1, f'=COUNTIF(C{n}:{get_column_letter(len(ids) + 2)}{n},"?*")')
    style_sheet(ws, [30, 38] + [7] * len(ids) + [10], freeze="C2")
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=len(ids) + 2):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="top")
            if cell.value:
                cell.fill = SUB_FILL
    leg = ws.max_row + 2
    ws.cell(leg, 1, "Legend").font = Font(name=ARIAL, size=10, bold=True)
    ws.cell(leg, 2, "F = field filter · D = date-range filter · GR = group rows · GC = group columns · C = detail column · "
                    "CΣ = column with summary · SF = custom summary formula · S = report sort").font = Font(name=ARIAL, size=10)

    # ------------------------------------------------------------- 6. Dashboards
    ws = wb.create_sheet("Dashboards")
    ws.append(["Dashboard ID", "Title", "API name", "Folder", "Description", "Dashboard type", "Running user (View as)",
               "Grid columns", "Row height (px)", "Components", "Theme / palette", "Dashboard filter", "Filter field",
               "Filter options (operator: equals)", "Shown on Home page"])
    for d in dashboards:
        for i, (fname, values, _ops) in enumerate(d["filters"]):
            field = "Opportunity Owner (FULL_NAME)" if "Owner" in fname else "Account: Industry (INDUSTRY)"
            ws.append([d["id"], d["title"], d["api"], "ABA Management Dashboards", d["description"] if i == 0 else "",
                       "Run as a specified user" if d["type"] == "SpecifiedUser" else d["type"], d["running_user"],
                       int(d["grid_cols"]), int(d["row_height"]), len(d["components"]), f"{d['theme']} / {d['palette']}",
                       fname, field, "All, " + ", ".join(values), "Yes – top region"])
    style_sheet(ws, [11, 34, 34, 24, 44, 20, 36, 9, 9, 11, 14, 18, 28, 60, 16], freeze="C2", table_name="DashboardSettings")

    # ------------------------------------------------------------- 7. Dashboard Components
    ws = wb.create_sheet("Dashboard Components")
    ws.append(["Dashboard ID", "Dashboard", "Component #", "Header", "Component type", "Grid row", "Grid column",
               "Row span", "Column span", "Source report ID", "Source report", "Source report filters", "Measure",
               "Groupings (X-axis → series)", "Sort", "Legend / labels", "Gauge or metric bands", "Table columns",
               "Table sort", "Responds to dashboard filters", "Drill-down", "Current value (formula)"])
    for d in dashboards:
        for i, c in enumerate(d["components"], 1):
            rep = by_api[c["report"]]
            rfilters = [OPEN_PARAM.get(rep["params"].get("open")), f"{label(rep['date_col'])} = {INTERVALS[rep['interval']][0]}"]
            rfilters += [f"{label(col)} {OPERATORS.get(o, o)} \"{v}\"" for col, o, v in rep["criteria"]]
            agg, col = c["summary"]
            measure = (f"{agg} of {label(col)}" if agg else label(col)) if col else "Table columns (see right)"
            gmin, gmax, bp1, bp2 = c["gauge"]
            if c["type"] == "Gauge":
                bands = f"0 – {int(bp1):,} red · {int(bp1):,} – {int(bp2):,} amber · {int(bp2):,} – {int(gmax):,} green (target {int(gmax):,})"
            elif c["type"] == "Metric":
                bands = (f"Red below {int(bp1):,}, amber to {int(bp2):,}, green above" if bp1 else "Single neutral colour (no target)")
            else:
                bands = "–"
            legend = ", ".join(x for x in [f"legend {c['legend']}" if c["legend"] else "",
                                            f"values {c['show_values']}" if c["show_values"] else "",
                                            f"total {c['show_total']}" if c["show_total"] else "",
                                            f"percent {c['show_pct']}" if c["show_pct"] else "",
                                            f"label '{c['metric_label']}'" if c["metric_label"] else ""] if x) or "–"
            tcols = ", ".join(f"{label(t[0])} ({t[1]})" for t in c["table_cols"]) or "–"
            tsort = (f"{label(c['table_sort'][0])} {'ascending' if c['table_sort'][1] == '1' else 'descending'}"
                     if c["table_sort"] else "–")
            ws.append([d["id"], d["title"], i, c["header"], c["type"], c["row"], c["col"], c["rows"], c["cols"],
                       rep["id"], rep["name"], "\n".join(rfilters), measure,
                       " → ".join(label(g) for g in c["groupings"]) or "–", c["sort_by"] or "–", legend, bands, tcols, tsort,
                       ", ".join(label(f) for f in c["filter_cols"]), "Enabled" if c["drill"] == "true" else "Off (opens report via View Report)",
                       None])
    style_sheet(ws, [11, 30, 10, 34, 14, 8, 9, 8, 8, 10, 32, 46, 24, 40, 20, 34, 46, 50, 22, 30, 18, 18],
                freeze="E2", table_name="DashboardComponents")
    comp_ws = ws

    # ------------------------------------------------------------- 8. Home Page
    ws = wb.create_sheet("Home Page")
    ws.append(["Lightning page", "Template", "Assigned to", "Region", "Position", "Component", "Shows", "Settings"])
    dash_title = {d["api"]: f"{d['id']} – {d['title']}" for d in dashboards}
    for region, pos, comp, props in home_rows:
        if comp == "desktopDashboards:embeddedDashboard":
            shows = "Embedded dashboard: " + dash_title[props["dashboardName"]]
            settings = f"Height {props.get('height')} px; hide on error {props.get('hideOnError')}"
        else:
            rep = rep_by_dev[props["reportName"]]
            shows = f"Report chart: {rep['id']} {rep['name']}"
            settings = f"Label '{props.get('label')}'; cache {props.get('cacheAge')} min; refresh button {props.get('showRefreshButton')}"
        ws.append([f"{home_label} (ABA_Sales_Home_Page)", home_template, "Sales app (standard__LightningSales), desktop Home tab",
                   region, pos, comp, shows, settings])
    style_sheet(ws, [30, 22, 34, 14, 9, 36, 52, 46], freeze="D2", table_name="HomePage")

    # ------------------------------------------------------------- 9. Opportunity Data
    ws = wb.create_sheet("Opportunity Data")
    base = ["Opportunity ID", "Opportunity Name", "Account Name", "Account: Industry", "Account: Billing Country",
            "Account: Segment", "Opportunity Owner", "Stage", "Closed", "Won", "Probability (%)", "Amount",
            "Expected Revenue", "Close Date", "Close Quarter", "Fiscal Year", "Forecast Category", "Type", "Lead Source",
            "Primary Campaign Source", "Offering", "Proposal Submission Date", "Win/Loss Reason", "Win/Loss Details",
            "Deferred Date", "Next Step", "Last Activity", "Created Date", "Age (days)", "Source of record"]
    ws.append(base + [f"In {i}" for i in ids])
    aba_accounts = {"Gulf Petrochem Services", "Najd Retail Group", "Red Sea Logistics", "Al Noor Healthcare", "Emirates Facility Mgmt",
                    "Desert Bloom Foods", "Riyadh Smart Utilities", "Khaleej Finance House", "Taiba Education Co",
                    "Oasis Manufacturing", "Bahrain Telecom Partners", "Qassim Agri Ventures"}
    to_date = lambda s: dt.date.fromisoformat(s[:10]) if s else None
    col = {h: i + 1 for i, h in enumerate(base)}
    L = lambda h: get_column_letter(col[h])
    for o in opps:
        acc = o.get("Account") or {}
        n = ws.max_row + 1
        row = [o["Id"], o["Name"], acc.get("Name"), acc.get("Industry"), acc.get("BillingCountry"), acc.get("Segment__c"),
               (o.get("Owner") or {}).get("Name"), o["StageName"], 1 if o["IsClosed"] else 0, 1 if o["IsWon"] else 0,
               (o["Probability"] or 0) / 100, o["Amount"], f"={L('Amount')}{n}*{L('Probability (%)')}{n}",
               to_date(o["CloseDate"]), o["Close_Quarter__c"], str(o["FiscalYear"]), o["ForecastCategoryName"], o["Type"],
               o["LeadSource"], (o.get("Campaign") or {}).get("Name"), o["Offering__c"], to_date(o["Proposal_Submission_Date__c"]),
               o["Win_Loss_Reason__c"], o["Win_Loss_Details__c"], to_date(o["Deferred_Date__c"]), o["NextStep"],
               to_date(o["LastActivityDate"]), to_date(o["CreatedDate"]),
               f"=IF({L('Closed')}{n}=1,{L('Close Date')}{n}-{L('Created Date')}{n},{as_of_ref}-{L('Created Date')}{n})",
               "ABA sample CSV" if acc.get("Name") in aba_accounts else "Created manually in the org (not part of the build)"]
        row += [1 if o["Id"] in member[i] else 0 for i in ids]
        ws.append(row)
    style_sheet(ws, [20, 44, 24, 18, 18, 12, 22, 18, 7, 6, 10, 13, 13, 12, 8, 8, 14, 32, 18, 16, 18, 13, 28, 30, 12, 16, 12, 12, 9, 22]
                + [7] * len(ids), freeze="C2", table_name="OpportunityData")
    last = ws.max_row
    for r in range(2, last + 1):
        ws[f"{L('Probability (%)')}{r}"].number_format = "0%"
        for h in ("Amount", "Expected Revenue"):
            ws[f"{L(h)}{r}"].number_format = "#,##0;(#,##0);-"
        for h in ("Close Date", "Proposal Submission Date", "Deferred Date", "Last Activity", "Created Date"):
            ws[f"{L(h)}{r}"].number_format = "dd-mmm-yyyy"
        ws[f"{L('Age (days)')}{r}"].number_format = "0"
        for j in range(len(base) + 1, len(base) + len(ids) + 1):
            c = ws.cell(r, j)
            c.alignment = Alignment(horizontal="center", vertical="top")
            if c.value == 1:
                c.fill = SUB_FILL
    flag_col = {i: get_column_letter(len(base) + k + 1) for k, i in enumerate(ids)}
    rng = lambda h: f"'Opportunity Data'!${L(h)}$2:${L(h)}${last}"
    frng = lambda i: f"'Opportunity Data'!${flag_col[i]}$2:${flag_col[i]}${last}"

    # ------------------------------------------------------------- 10. Report Results
    ws = wb.create_sheet("Report Results")
    ws.append(["Report ID", "Report name", "Records (formula)", "Sum of Amount", "Sum of Expected Revenue", "Won records",
               "Closed records", "Win rate", "Average Amount", "Average Age (days)", "Live count from Salesforce",
               "Matches Salesforce?", "Note"])
    res_row = {}
    for r in reports:
        i = r["id"]
        n = ws.max_row + 1
        res_row[i] = n
        lc = live.get(i, ("", None))[1]
        ws.append([i, r["name"], f"=SUM({frng(i)})", f"=SUMPRODUCT({frng(i)},{rng('Amount')})",
                   f"=SUMPRODUCT({frng(i)},{rng('Expected Revenue')})", f"=SUMPRODUCT({frng(i)},{rng('Won')})",
                   f"=SUMPRODUCT({frng(i)},{rng('Closed')})", f"=IF(G{n}=0,\"-\",F{n}/G{n})", f"=IF(C{n}=0,\"-\",D{n}/C{n})",
                   f"=IF(C{n}=0,\"-\",SUMPRODUCT({frng(i)},{rng('Age (days)')})/C{n})", lc,
                   f"=IF(C{n}=K{n},\"Yes\",\"Check\")",
                   "Win rate is the Win Rate % formula used by D1 and D2." if i == "R09" else
                   "Average Amount and Average Age are the report's summaries." if i == "R18" else ""])
    style_sheet(ws, [9, 38, 11, 15, 15, 9, 9, 9, 14, 12, 12, 11, 40], freeze="C2", table_name="ReportResults")
    for r in range(2, ws.max_row + 1):
        for c in ("D", "E", "I"):
            ws[f"{c}{r}"].number_format = "#,##0;(#,##0);-"
        ws[f"H{r}"].number_format = "0%"
        ws[f"J{r}"].number_format = "0.0"
    note = ws.max_row + 2
    ws.cell(note, 1, "How it works").font = Font(name=ARIAL, size=10, bold=True)
    ws.cell(note, 2, "Each figure is calculated from the 1/0 'In Rxx' flags on the Opportunity Data sheet. A flag is 1 when "
                     "the opportunity meets the report's filters, read from Salesforce on the as-of date. "
                     "'Live count from Salesforce' is the record count returned by running the report itself.").font = Font(name=ARIAL, size=10)
    ws.cell(note, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=note, start_column=2, end_row=note, end_column=8)
    ws.row_dimensions[note].height = 45

    # current value formulas on Dashboard Components
    for r in range(2, comp_ws.max_row + 1):
        rid, ctype, measure = comp_ws[f"J{r}"].value, comp_ws[f"E{r}"].value, comp_ws[f"M{r}"].value
        n = res_row[rid]
        if "Win Rate" in str(measure):
            comp_ws[f"V{r}"] = f"='Report Results'!H{n}"
            comp_ws[f"V{r}"].number_format = "0%"
        elif ctype == "FlexTable":
            comp_ws[f"V{r}"] = f"='Report Results'!C{n}&\" rows\""
        else:
            comp_ws[f"V{r}"] = f"='Report Results'!D{n}"
            comp_ws[f"V{r}"].number_format = "#,##0"
        comp_ws[f"V{r}"].font = Font(name=ARIAL, size=10, color="008000")

    # ------------------------------------------------------------- 11. Folders & Access
    ws = wb.create_sheet("Folders & Access")
    ws.append(["Folder", "API name", "Folder type", "Shared with", "Contents"])
    for api, kind in (("ABA_Management_Reports", "Report folder"), ("ABA_Sales_Reports", "Report folder"),
                      ("ABA_Management_Dashboards", "Dashboard folder")):
        if kind == "Report folder":
            contents = ", ".join(f"{r['id']} {r['name']}" for r in reports if r["folder"] == api)
        else:
            contents = ", ".join(f"{d['id']} {d['title']}" for d in dashboards)
        ws.append([api.replace("_", " "), api, kind, FOLDER_SHARING[api], contents])
    ws.append([])
    ws.append(["Dashboard viewing", "", "", "Both dashboards run as Suraj OD (System Administrator), so every viewer sees all opportunities.", ""])
    style_sheet(ws, [28, 28, 16, 44, 90], freeze=None)

    for sheet in wb.worksheets:
        sheet.sheet_view.zoomScale = 90
    wb.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
